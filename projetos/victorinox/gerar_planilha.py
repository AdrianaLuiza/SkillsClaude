"""Monta a planilha de valores do Projeto Victorinox com os logos no cabecalho."""
import pathlib
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AbsoluteAnchor, OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU

BASE = str(pathlib.Path(__file__).parent / "assets")
DEST = str(pathlib.Path(__file__).parent / "Projeto_Victorinox_Dez26_Valores.xlsx")

NAVY = "1B3A5C"
AZUL = "4A9FD8"
VINO = "6E0C24"
CINZA = "F2F6FA"
LINHA = "B8C6D6"

thin = Side(style="thin", color=LINHA)
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
ws = wb.active
ws.title = "Valores Dez-26"

# ----------------------------------------------------------------- larguras
widths = {"A": 26, "B": 5, "C": 6, "D": 6, "E": 10, "F": 9, "G": 9, "H": 9,
          "I": 9, "J": 9, "K": 10, "L": 10, "M": 10, "N": 10, "O": 12, "P": 11}
for col, w in widths.items():
    ws.column_dimensions[col].width = w
LAST = "P"

# ------------------------------------------------------------- cabecalho/logos
for r in range(1, 6):
    ws.row_dimensions[r].height = 17
ws.row_dimensions[6].height = 6

ws.merge_cells("E1:L2")
ws["E1"] = "PROJETO VICTORINOX"
ws["E1"].font = Font(name="Calibri", size=18, bold=True, color=NAVY)
ws["E1"].alignment = Alignment(horizontal="center", vertical="center")

ws.merge_cells("E3:L3")
ws["E3"] = "Proposta comercial de hospedagem  |  Valor Pacote  |  Dez/26"
ws["E3"].font = Font(name="Calibri", size=11, color="44546A")
ws["E3"].alignment = Alignment(horizontal="center", vertical="center")

ws.merge_cells("E4:L5")
ws["E4"] = ("Apresentado por: Pousada Camburí   ·   Responsável: Rodrigo Sulina\n"
            "Data: 25/08/2026   ·   Revisão: 00")
ws["E4"].font = Font(name="Calibri", size=9, color="44546A")
ws["E4"].alignment = Alignment(horizontal="center", vertical="center", wrapText=True)


def place(path, anchor_cell, px_w, px_h, dx=0, dy=0):
    img = XLImage(path)
    img.width, img.height = px_w, px_h
    col = ws[anchor_cell].column - 1
    row = ws[anchor_cell].row - 1
    marker = AnchorMarker(col=col, colOff=pixels_to_EMU(dx), row=row, rowOff=pixels_to_EMU(dy))
    img.anchor = OneCellAnchor(_from=marker,
                               ext=XDRPositiveSize2D(pixels_to_EMU(px_w), pixels_to_EMU(px_h)))
    ws.add_image(img)


# Pousada Camburí (empresa proponente) à esquerda; Victorinox (contratante) à direita
place(f"{BASE}/camburi.png", "A1", 216, 93, dx=8, dy=6)
place(f"{BASE}/victorinox.png", "N1", 128, 70, dx=54, dy=17)

# faixa de separacao
for c in range(1, 17):
    ws.cell(row=6, column=c).fill = PatternFill("solid", fgColor=NAVY)

# ------------------------------------------------------------------- tabela
GRUPOS = [
    ("Relação Apartamentos", "A", "B"),
    ("Capacidade", "C", "D"),
    ("Valores", "E", "H"),
    ("Refeição (04 a 06)", "I", "J"),
    ("Valor Pacote", "K", "N"),
    ("Diária adicional semana", "O", "P"),
]
SUB = ["Categoria", "Nº", "pax", "aux", "duplo", "3ª pax", "4ª pax", "pax add",
       "02 pax", "01 pax", "02 pax", "03 pax", "04 pax", "05 pax", "duplo", "extra"]

R_GRP, R_SUB, R_INI = 7, 8, 9
ws.row_dimensions[R_GRP].height = 20
ws.row_dimensions[R_SUB].height = 26

for titulo, c0, c1 in GRUPOS:
    ws.merge_cells(f"{c0}{R_GRP}:{c1}{R_GRP}")
    cel = ws[f"{c0}{R_GRP}"]
    cel.value = titulo
    cel.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    cel.alignment = Alignment(horizontal="center", vertical="center")
    for c in range(ws[f"{c0}1"].column, ws[f"{c1}1"].column + 1):
        cc = ws.cell(row=R_GRP, column=c)
        cc.fill = PatternFill("solid", fgColor=NAVY)
        cc.border = box

for i, nome in enumerate(SUB, start=1):
    cel = ws.cell(row=R_SUB, column=i, value=nome)
    cel.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    cel.alignment = Alignment(horizontal="center", vertical="center", wrapText=True)
    cel.fill = PatternFill("solid", fgColor=AZUL)
    cel.border = box

D = "-"
DADOS = [
    ("terreo duplo",            1, 2, 1, 1780, D,   D,   534, 520, 260, 2300, 3094, D,    D,    623, 267),
    ("terreo duplo varanda",    2, 2, 1, 1869, D,   D,   560, 520, 260, 2389, 3209, D,    D,    654, 280),
    ("terreo duplo",            3, 2, 1, 1780, D,   D,   534, 520, 260, 2300, 3094, D,    D,    623, 267),
    ("terreo duplo varanda",    4, 2, 1, 1869, D,   D,   560, 520, 260, 2389, 3209, D,    D,    654, 280),
    ("terreo duplo",            5, 2, D, 1780, D,   D,   D,   520, 260, 2300, D,    D,    D,    623, 267),
    ("terreo duplo varanda",    6, 2, D, 1869, D,   D,   D,   520, 260, 2389, D,    D,    D,    654, 280),
    ("superior triplo varanda", 7, 3, 1, 2055, 617, D,   617, 520, 260, 2575, 3452, 4329, D,    719, 308),
    ("superior triplo varanda", 8, 3, 1, 2055, 617, D,   617, 520, 260, 2575, 3452, 4329, D,    719, 308),
    ("superior duplo balcão",   9, 2, D, 1869, D,   D,   D,   520, 260, 2389, D,    D,    D,    654, 280),
    ("superior duplo balcão",  10, 2, D, 1869, D,   D,   D,   520, 260, 2389, D,    D,    D,    654, 280),
    ("superior duplo balcão",  11, 2, D, 1869, D,   D,   D,   520, 260, 2389, D,    D,    D,    654, 280),
    ("superior duplo balcão",  12, 2, D, 1869, D,   D,   D,   520, 260, 2389, D,    D,    D,    654, 280),
    ("superior quádruplo",     13, 4, 1, 2157, 647, 647, 647, 520, 260, 2677, 3584, 4491, 5398, 755, 324),
    ("superior quádruplo",     14, 4, 1, 2157, 647, 647, 647, 520, 260, 2677, 3584, 4491, 5398, 755, 324),
    ("superior triplo",        15, 3, 1, 2055, 617, D,   617, 520, 260, 2575, 3452, 4329, D,    719, 308),
    ("terreo triplo",          16, 3, D, 1820, 546, D,   546, 520, 260, 2340, 3146, D,    D,    637, 273),
    ("terreo triplo",          17, 3, D, 1820, 546, D,   546, 520, 260, 2340, 3146, D,    D,    637, 273),
    ("terreo duplo",           18, 2, D, 1780, D,   D,   D,   520, 260, 2300, D,    D,    D,    623, 267),
]

for i, linha in enumerate(DADOS):
    r = R_INI + i
    ws.row_dimensions[r].height = 16
    fundo = PatternFill("solid", fgColor=CINZA) if i % 2 else None
    for c, valor in enumerate(linha, start=1):
        cel = ws.cell(row=r, column=c, value=valor)
        cel.border = box
        cel.font = Font(name="Calibri", size=10)
        if fundo:
            cel.fill = fundo
        if c == 1:
            cel.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        else:
            cel.alignment = Alignment(horizontal="center", vertical="center")
        if isinstance(valor, int) and c >= 5:
            cel.number_format = "#,##0"
        if c in (11, 12, 13, 14) and isinstance(valor, int):
            cel.font = Font(name="Calibri", size=10, bold=True, color=NAVY)

R_TOT = R_INI + len(DADOS)
ws.row_dimensions[R_TOT].height = 18
ws.merge_cells(f"A{R_TOT}:B{R_TOT}")
ws[f"A{R_TOT}"] = "Capacidade total de pessoas: 54"
ws[f"C{R_TOT}"] = f"=SUM(C{R_INI}:C{R_TOT - 1})"
ws[f"D{R_TOT}"] = f"=COUNT(D{R_INI}:D{R_TOT - 1})"
for c in range(1, 17):
    cel = ws.cell(row=R_TOT, column=c)
    cel.fill = PatternFill("solid", fgColor=NAVY)
    cel.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    cel.border = box
    cel.alignment = Alignment(horizontal="center", vertical="center")
ws[f"A{R_TOT}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

# --------------------------------------------------------- valor negociado
R_NEG = R_TOT + 2
ws.merge_cells(f"A{R_NEG}:D{R_NEG}")
ws[f"A{R_NEG}"] = "VALOR NEGOCIADO — referências"
ws[f"A{R_NEG}"].font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ws[f"A{R_NEG}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
for c in range(1, 5):
    ws.cell(row=R_NEG, column=c).fill = PatternFill("solid", fgColor=VINO)
ws.row_dimensions[R_NEG].height = 20

NEG = [
    ("duplo",     2300, "02 adultos ou 01 adulto + 01 criança"),
    ("triplo",    3000, "02 adultos + 01 criança"),
    ("quádruplo", 3800, "02 adultos + 02 crianças"),
]
for i, (cat, val, desc) in enumerate(NEG):
    r = R_NEG + 1 + i
    ws.cell(row=r, column=1, value=cat).font = Font(name="Calibri", size=10, bold=True)
    ws.cell(row=r, column=1).alignment = Alignment(horizontal="left", vertical="center", indent=1)
    c2 = ws.cell(row=r, column=2, value=val)
    c2.number_format = "#,##0"
    c2.font = Font(name="Calibri", size=10, bold=True, color=VINO)
    c2.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    c3 = ws.cell(row=r, column=3, value=desc)
    c3.font = Font(name="Calibri", size=10)
    c3.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for c in range(1, 7):
        ws.cell(row=r, column=c).border = box

# ---------------------------------------------------------------- observacoes
R_OBS = R_NEG + 5
ws.merge_cells(f"A{R_OBS}:{LAST}{R_OBS}")
ws[f"A{R_OBS}"] = "OBSERVAÇÕES"
ws[f"A{R_OBS}"].font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ws[f"A{R_OBS}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
for c in range(1, 17):
    ws.cell(row=R_OBS, column=c).fill = PatternFill("solid", fgColor=NAVY)
ws.row_dimensions[R_OBS].height = 20

OBS = [
    "Diária de semana extra não contempla alimentação; mais favorável prato à la carte (mais econômico).",
    "Itens de consumo como bebidas serão lançados individualmente nos apartamentos, conforme pedidos.",
    "Os apartamentos são mais amplos conforme a capacidade de acomodação de cada categoria (duplo, triplo ou quádruplo).",
    "Pagamentos podem ser realizados por link de pagamento em até 12x sem juros, ou com desconto de 5% à vista via PIX.",
    "Os interessados deverão enviar dados de cadastro antecipadamente, em conformidade com a LGPD (Lei Geral de Proteção de Dados).",
]
for i, texto in enumerate(OBS):
    r = R_OBS + 1 + i
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=16)
    cel = ws.cell(row=r, column=1, value=f"•  {texto}")
    cel.font = Font(name="Calibri", size=9.5, color="44546A")
    cel.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[r].height = 15

# ---------------------------------------------------------------- acabamento
ws.freeze_panes = f"A{R_INI}"
ws.sheet_view.showGridLines = False
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_area = f"A1:{LAST}{R_OBS + len(OBS)}"

wb.save(DEST)
print("salvo:", DEST)
